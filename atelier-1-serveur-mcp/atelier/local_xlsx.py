"""Backend de stockage — formulaires NDF en fichiers .xlsx locaux (openpyxl).

⚙️  PLOMBERIE — ne pas toucher pendant l'atelier.

Copie les vrais templates (templates/<ENV>.xlsx) dans le dossier de travail et écrit
dedans : identité (B10/B11) + lignes de factures (A/B, 15→28). Persistant sur disque.

Sortie : $NDF_DIR (défaut ./notes_de_frais), un sous-dossier par enveloppe (CE / ACADEMY).
Dépôt : $INBOX_DIR (défaut ./factures_a_classer). Total B29 garde sa formule SUBTOTAL ;
balance() resomme B15:B28 côté Python.
"""

import datetime
import os
import re
import shutil
import unicodedata
from pathlib import Path

import openpyxl

from config import ENVELOPES, FORM

ROOT = Path(__file__).resolve().parent.parent
TEMPLATES_DIR = ROOT / "templates"
BASE_DIR = Path(os.environ.get("NDF_DIR", ROOT / "notes_de_frais"))
# Dossier "boîte de dépôt" : l'utilisateur y glisse les factures (PDF) à classer.
INBOX_DIR = Path(os.environ.get("INBOX_DIR", ROOT / "factures_a_classer"))
_INVOICE_EXT = {".pdf"}

# Sous-dossier + suffixe de fichier par enveloppe (APPRENTISSAGE = "Académie" → ACADEMY).
_SUBDIR = {"CE": "CE", "APPRENTISSAGE": "ACADEMY"}
_FILE_SUFFIX = {"CE": "NDF_CE", "APPRENTISSAGE": "NDF_RembApp"}


def _env_dir(envelope: str) -> Path:
    """Dossier d'une enveloppe : <NDF_DIR>/<CE|ACADEMY>."""
    return BASE_DIR / _SUBDIR[envelope.upper()]


def _norm_nom(nom: str) -> str:
    """NOM en MAJUSCULES, sans accents ni séparateurs (pour le nom de fichier)."""
    n = unicodedata.normalize("NFKD", nom).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Za-z0-9]+", "", n).upper()


def _filename(envelope: str, nom: str) -> str:
    """Nom du fichier : aaaammjj_NOM_<suffixe>.xlsx (date du jour + nom en majuscules)."""
    today = datetime.date.today().strftime("%Y%m%d")
    return f"{today}_{_norm_nom(nom)}_{_FILE_SUFFIX[envelope.upper()]}.xlsx"


def _existing(envelope: str) -> Path | None:
    """Le .xlsx déjà créé pour cette enveloppe, ou None (ne crée rien)."""
    d = _env_dir(envelope)
    if not d.is_dir():
        return None
    hits = sorted(d.glob(f"*_{_FILE_SUFFIX[envelope.upper()]}.xlsx"))
    return hits[0] if hits else None


def _open(path: Path):
    wb = openpyxl.load_workbook(path)
    tab = FORM["tab"]
    ws = wb[tab] if tab in wb.sheetnames else wb.active
    return wb, ws


def _get_or_create(envelope: str, nom: str) -> Path:
    """Le formulaire de l'enveloppe ; le crée depuis le template (nommé avec NOM) si absent."""
    p = _existing(envelope)
    if p is not None:
        return p
    env = envelope.upper()
    tpl = TEMPLATES_DIR / f"{env}.xlsx"
    if not tpl.exists():
        raise FileNotFoundError(f"Template manquant : {tpl}")
    (_env_dir(env) / "factures").mkdir(parents=True, exist_ok=True)
    p = _env_dir(env) / _filename(env, nom)
    shutil.copyfile(tpl, p)
    return p


def record_line(prenom: str, nom: str, envelope: str, descriptif: str, montant: float) -> tuple[int, str]:
    """Écrit l'identité (B10/B11) + une ligne de facture sur la prochaine ligne libre.

    Crée le formulaire au besoin. Renvoie (numéro de ligne, chemin du .xlsx).
    """
    p = _get_or_create(envelope, nom)
    wb, ws = _open(p)
    ws[FORM["nom_cell"]] = nom
    ws[FORM["prenom_cell"]] = prenom
    first, last = FORM["first_line_row"], FORM["last_line_row"]
    dcol, acol = FORM["desc_col"], FORM["amount_col"]
    row = None
    for r in range(first, last + 1):
        if not ws[f"{dcol}{r}"].value:
            row = r
            break
    if row is None:
        raise RuntimeError(f"Formulaire {envelope} plein (14 lignes max)")
    ws[f"{dcol}{row}"] = descriptif
    ws[f"{acol}{row}"] = float(montant)
    wb.save(p)
    return row, str(p)


def factures_dir(envelope: str) -> str:
    """Dossier où l'utilisateur dépose lui-même les justificatifs (créé au besoin)."""
    d = _env_dir(envelope) / "factures"
    d.mkdir(parents=True, exist_ok=True)
    return str(d)


def list_inbox() -> list[dict]:
    """Factures déposées dans la boîte à classer ($INBOX_DIR) — créée au besoin."""
    INBOX_DIR.mkdir(parents=True, exist_ok=True)
    return [
        {"filename": f.name, "path": str(f)}
        for f in sorted(INBOX_DIR.iterdir())
        if f.is_file() and f.suffix.lower() in _INVOICE_EXT
    ]


def archive(envelope: str, src_path: str) -> str:
    """Déplace l'original de la facture dans <enveloppe>/factures/ (gère les collisions)."""
    src = Path(src_path).expanduser()
    if not src.is_file():
        raise FileNotFoundError(f"Fichier introuvable : {src_path}")
    dest_dir = _env_dir(envelope) / "factures"
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / src.name
    i = 1
    while dest.exists():
        dest = dest_dir / f"{src.stem}_{i}{src.suffix}"
        i += 1
    shutil.move(str(src), str(dest))
    return str(dest)


def balance(envelope: str) -> dict:
    """Solde d'une enveloppe — LECTURE SEULE (ne crée jamais de formulaire)."""
    env_key = envelope.upper()
    env = ENVELOPES[env_key]
    rate, cap = env["reimbursement_rate"], env["cap_reimbursed"]
    p = _existing(env_key)
    if p is None:
        return {"envelope": env_key, "note": "aucune dépense enregistrée pour le moment",
                "engaged": 0.0, "reimbursed": 0.0, "rate": rate,
                "cap_reimbursed": cap, "remaining_reimbursed": cap, "over_cap": False}
    _, ws = _open(p)
    first, last = FORM["first_line_row"], FORM["last_line_row"]
    acol = FORM["amount_col"]
    engaged = 0.0
    for r in range(first, last + 1):
        try:
            engaged += float(ws[f"{acol}{r}"].value)
        except (TypeError, ValueError):
            pass
    reimbursed = engaged * rate
    remaining = max(0.0, cap - reimbursed)
    return {
        "envelope": env_key,
        "engaged": round(engaged, 2),
        "reimbursed": round(reimbursed, 2),
        "rate": rate,
        "cap_reimbursed": cap,
        "remaining_reimbursed": round(remaining, 2),
        "remaining_engageable": round(remaining / rate, 2) if rate else 0.0,
        "over_cap": reimbursed > cap,
        "local": True,
        "path": str(p),
    }
