"""Backend de stockage — formulaires NDF en fichiers .xlsx locaux (openpyxl).

Copie les vrais templates (templates/<ENV>.xlsx) dans un dossier de travail local et
écrit dedans : identité (B10/B11) + lignes de factures (A/B, 15→28). Persistant sur
disque → on peut ouvrir le .xlsx dans Excel/LibreOffice.

La cellule total B29 garde sa formule SUBTOTAL (jamais touchée) ; balance() recalcule
l'engagé en sommant B15:B28 côté Python (openpyxl n'évalue pas les formules).

Sortie : $LOCAL_XLSX_DIR (défaut ./local_data), un sous-dossier par user_id.
"""

import os
import re
import shutil
from pathlib import Path

import openpyxl

from src.config import ENVELOPES, FORM

ROOT = Path(__file__).resolve().parent.parent
TEMPLATES_DIR = ROOT / "templates"
BASE_DIR = Path(os.environ.get("LOCAL_XLSX_DIR", ROOT / "local_data"))


def _safe(user_id: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", user_id) or "user"


def _user_dir(user_id: str) -> Path:
    d = BASE_DIR / _safe(user_id)
    d.mkdir(parents=True, exist_ok=True)
    return d


def _path(user_id: str, envelope: str) -> Path:
    return _user_dir(user_id) / f"{envelope.upper()}.xlsx"


def _ensure_file(user_id: str, envelope: str) -> Path:
    """Copie le template de l'enveloppe dans le dossier du user si absent."""
    env = envelope.upper()
    p = _path(user_id, env)
    if not p.exists():
        tpl = TEMPLATES_DIR / f"{env}.xlsx"
        if not tpl.exists():
            raise FileNotFoundError(f"Template manquant : {tpl}")
        shutil.copyfile(tpl, p)
    return p


def _open(path: Path):
    wb = openpyxl.load_workbook(path)
    tab = FORM["tab"]
    ws = wb[tab] if tab in wb.sheetnames else wb.active
    return wb, ws


def setup_user(user_id: str, nom: str = "", prenom: str = "") -> dict:
    sheets = {}
    for env in ENVELOPES:
        p = _ensure_file(user_id, env)
        sheets[env] = {"id": str(p), "path": str(p)}
    if nom or prenom:
        for env in ENVELOPES:
            set_identity(user_id, env, nom, prenom)
    return {"folder_id": str(_user_dir(user_id)), "sheets": sheets, "local": True}


def set_identity(user_id: str, envelope: str, nom: str, prenom: str) -> None:
    p = _ensure_file(user_id, envelope)
    wb, ws = _open(p)
    ws[FORM["nom_cell"]] = nom
    ws[FORM["prenom_cell"]] = prenom
    wb.save(p)


def add_invoice_line(user_id: str, envelope: str, descriptif: str, montant: float) -> int:
    p = _ensure_file(user_id, envelope)
    wb, ws = _open(p)
    first, last = FORM["first_line_row"], FORM["last_line_row"]
    dcol, acol = FORM["desc_col"], FORM["amount_col"]
    row = None
    for r in range(first, last + 1):
        if not ws[f"{dcol}{r}"].value:
            row = r
            break
    if row is None:
        raise RuntimeError(f"Formulaire {envelope} plein (14 lignes max)")
    ws[f"{dcol}{row}"] = descriptif
    ws[f"{acol}{row}"] = float(montant)
    wb.save(p)
    return row


def balance(user_id: str, envelope: str) -> dict:
    env_key = envelope.upper()
    env = ENVELOPES[env_key]
    p = _ensure_file(user_id, env_key)
    _, ws = _open(p)
    first, last = FORM["first_line_row"], FORM["last_line_row"]
    acol = FORM["amount_col"]
    engaged = 0.0
    for r in range(first, last + 1):
        try:
            engaged += float(ws[f"{acol}{r}"].value)
        except (TypeError, ValueError):
            pass
    rate, cap = env["reimbursement_rate"], env["cap_reimbursed"]
    reimbursed = engaged * rate
    remaining = max(0.0, cap - reimbursed)
    return {
        "envelope": env_key,
        "engaged": round(engaged, 2),
        "reimbursed": round(reimbursed, 2),
        "rate": rate,
        "cap_reimbursed": cap,
        "remaining_reimbursed": round(remaining, 2),
        "remaining_engageable": round(remaining / rate, 2) if rate else 0.0,
        "local": True,
        "path": str(p),
    }
